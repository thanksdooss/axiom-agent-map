#!/usr/bin/env python3
"""더미 데이터 v2 교차 검증. build_dataset.py와 별개로 원본 파일을 다시 읽어 정답과 맞춰 봅니다.

실행: python3 Data/v2/tools/verify_dataset.py   (표준 라이브러리만 필수, PyYAML·openpyxl이 있으면 추가 검사)
"""
import csv
import json
import re
import shutil
import subprocess
import sys
import zipfile
from collections import defaultdict
from html import unescape
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "raw"
problems = []


def check(cond, msg):
    if not cond:
        problems.append(msg)


vocab = json.loads((ROOT / "vocabulary/vocabulary.json").read_text(encoding="utf-8"))
key = json.loads((ROOT / "answer_key/extraction_answer_key.json").read_text(encoding="utf-8"))
gold = json.loads((ROOT / "answer_key/query_golden_set.json").read_text(encoding="utf-8"))
rels = key["relationships"]

# ---------------------------------------------------------------- 1. raw 파일 독립 파싱
raw_json = json.loads((RAW / "01_영업_마케팅.json").read_text(encoding="utf-8"))
yaml_lines = (RAW / "02_개발_IT운영.yaml").read_text(encoding="utf-8").splitlines()
csv_rows = list(csv.reader((RAW / "03_고객지원_구매.csv").read_text(encoding="utf-8").splitlines()))

z = zipfile.ZipFile(RAW / "04_인사_재무.xlsx")
xlsx = {}
for ref, inner in re.findall(r'<c r="([A-Z]+\d+)"[^>]*?(?:/>|>(.*?)</c>)', z.read("xl/worksheets/sheet1.xml").decode(), re.S):
    t = re.search(r"<t[^>]*>(.*?)</t>", inner or "", re.S)
    v = re.search(r"<v>(.*?)</v>", inner or "", re.S)
    xlsx[ref] = unescape(t.group(1)) if t else (v.group(1) if v else "")

doc = zipfile.ZipFile(RAW / "05_법무_경영기획.docx").read("word/document.xml").decode()
paras = [unescape("".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", p, re.S))) for p in re.findall(r"<w:p>.*?</w:p>", doc, re.S)]


def locate(file, loc):
    """출처 위치에 있는 원문 텍스트."""
    if file.endswith(".json"):
        node = raw_json
        for part in loc.strip("/").split("/"):
            node = node[int(part)] if isinstance(node, list) else node[part]
        return node if isinstance(node, str) else ("" if node is None else str(node))
    if file.endswith(".yaml"):
        return yaml_lines[int(re.match(r"(\d+)행", loc)[1]) - 1]
    if file.endswith(".csv"):
        m = re.match(r"(\d+)행 (\w+)열", loc)
        row = csv_rows[int(m[1]) - 1]
        return row[csv_rows[0].index(m[2])]
    if file.endswith(".xlsx"):
        return xlsx.get(loc.split("!")[1], "")
    if file.endswith(".docx"):
        m = re.match(r"문단 (\d+)(?: 문장 (\d+))?", loc)
        p = paras[int(m[1]) - 1]
        if not m[2]:
            return p
        return re.split(r"(?<=\.)\s", p)[int(m[2]) - 1]
    return None


# ---------------------------------------------------------------- 2. 근거 위치 검사
n_ev = 0
for r in rels:
    for e in r["evidence"]:
        if not e["file"].startswith(("01", "02", "03", "04", "05")):
            continue
        n_ev += 1
        found = locate(e["file"], e["locator"])
        text = e["text"] or ""
        check(found is not None and text.strip("'") in found, f"근거 불일치 {r['id']}: {e['locator']} → {found!r} vs {text!r}")
for c in key["special_cases_in_files"] + [dict(evidence=p["evidence"]) for p in key["pending_entities"]]:
    e = c["evidence"]
    found = locate(e["file"], e["locator"])
    check(found is not None and (e["text"] or "").strip("'") in found, f"특수 사례 근거 불일치 {e['locator']}")

# ---------------------------------------------------------------- 3. 표준화 다시 계산 (어휘 파일만 사용)
alias = defaultdict(dict)
for kind, items in [("department", vocab["departments"]), ("work_type", vocab["work_types"]),
                    ("platform", vocab["platforms"]), ("data", vocab["data_assets"])]:
    for it in items:
        for label in [it["prefLabel"], *it.get("altLabel", [])]:
            alias[kind][label] = it["id"]
review = vocab["needs_review_labels"]
unreg = vocab["unregistered_labels"]
agent_ids = {re.sub(r"\s+", "", a["name"]): a["id"] for a in key["agents"] + key["sample_agents"]}


def norm(kind, raw):
    if raw in alias[kind]:
        return alias[kind][raw], "CONFIRMED"
    if raw in review.get(kind, {}):
        return review[kind][raw]["candidate"], "CANDIDATE"
    if raw in unreg.get(kind, {}):
        return None, "PENDING"
    raise KeyError(f"{kind} 미정의 표기 {raw!r}")


expected = set()
for a in key["agents"]:
    raw = a["raw"]
    aid = a["id"]
    if raw["dept"]:
        d, s = norm("department", raw["dept"])
    else:
        d, s = next(f["candidate"] for f in a["review"] if f["field"] == "department"), "CANDIDATE"
    expected.add(("BELONGS_TO", aid, d, s))
    expected.add(("HAS_WORK_TYPE", aid, norm("work_type", raw["type"])[0], "CONFIRMED"))
    for p in raw["platforms"]:
        expected.add(("USES_PLATFORM", aid, *norm("platform", p)))
    for x in raw["data"]:
        t, s = norm("data", x)
        if t:
            expected.add(("ACCESSES_DATA", aid, t, s))
    for c in raw["calls"]:
        k = re.sub(r"\s+", "", c)
        if c in unreg["agent"]:
            continue
        if k not in agent_ids:  # 모호한 이름: 정답 세트의 선택 결과 사용
            chosen = next(f["candidate"] for f in a["review"] if f["field"] == "calls" and f["raw"] == c)
            expected.add(("USES_RESULT_OF", aid, chosen, "CANDIDATE"))
            continue
        if agent_ids[k] != aid:
            expected.add(("USES_RESULT_OF", aid, agent_ids[k], "CONFIRMED"))
# 제공하는 쪽이 주어인 Word 문장(A044) — 근거 메모가 달린 관계
for r in rels:
    for e in r["evidence"]:
        if e.get("direction_note"):
            expected.add((r["type"], r["source"], r["target"], r["status"]))
sample = {a["id"] for a in key["sample_agents"]}
actual = {(r["type"], r["source"], r["target"], r["status"]) for r in rels
          if r["type"] != "STORED_IN" and r["source"] not in sample}
for x in sorted(expected - actual):
    problems.append(f"정답 세트에 빠진 관계 {x}")
for x in sorted(actual - expected):
    problems.append(f"원본 근거로 설명되지 않는 관계 {x}")

# Word 원문 토큰이 실제 문장에 있는지
docx_agents = [a for a in key["agents"] if a["source_file"].endswith(".docx")]
doc_text = "\n".join(paras)
for a in docx_agents:
    for v in [a["raw"]["dept"], a["raw"]["type"], *a["raw"]["platforms"], *a["raw"]["data"], *a["raw"]["calls"]]:
        check(v in doc_text, f"Word 원문에 없는 표기 {a['id']} {v!r}")

# ---------------------------------------------------------------- 4. 질의 정답 다시 계산 (집합 확장 방식)
seed = list(csv.DictReader((ROOT / "seed/relationships.csv").read_text(encoding="utf-8-sig").splitlines()))
name = {}
for r in seed:
    name[r["source_id"]], name[r["target_id"]] = r["source_name"], r["target_name"]


def edges(scope):
    ok = {"all_statuses": {"SAMPLE", "CANDIDATE", "CONFIRMED"}, "confirmed_only": {"CONFIRMED"}}[scope]
    return [r for r in seed if r["status"] in ok]


def impact(anchor, scope):
    es = edges(scope)
    first = {"A": "USES_RESULT_OF", "D": "ACCESSES_DATA", "P": "USES_PLATFORM"}[anchor[0]]
    depth = {anchor: 0}
    frontier, level = {anchor}, 0
    while frontier:
        level += 1
        t = first if level == 1 else "USES_RESULT_OF"
        nxt = {r["source_id"] for r in es if r["type"] == t and r["target_id"] in frontier} - depth.keys()
        for n in nxt:
            depth[n] = level
        frontier = nxt
    depth.pop(anchor)
    return {name[k]: v for k, v in depth.items()}


def from_gold(exp):
    out = {n: 1 for n in exp["direct"]}
    out.update({x["name"]: x["depth"] for x in exp["indirect"]})
    return out


gq = {q["id"]: q for q in gold["queries"]}
for qid, anchor in [("Q04", "A001"), ("Q15", "D001"), ("Q16", "A005"), ("Q19", "P01")]:
    for scope in ("all_statuses", "confirmed_only"):
        check(impact(anchor, scope) == from_gold(gq[qid]["expected"][scope]), f"{qid} {scope} 영향 탐색 불일치")
sens = {d["id"]: d["sensitivity"] for d in vocab["data_assets"]}
for scope in ("all_statuses", "confirmed_only"):
    hi = {r["source_id"] for r in edges(scope) if r["type"] == "ACCESSES_DATA" and sens.get(r["target_id"]) == "상"}
    check(len(hi) == gq["Q10"]["expected"][scope]["count"], f"Q10 {scope} 개수 불일치 {len(hi)}")
calls = defaultdict(set)
for r in edges("all_statuses"):
    if r["type"] == "USES_RESULT_OF":
        calls[r["source_id"]].add(r["target_id"])
cyc = sorted({tuple(sorted((a, b))) for a in calls for b in calls[a] if a in calls.get(b, set())})
check([[name[a], name[b]] for a, b in cyc] == gq["Q18"]["expected"]["all_statuses"], "Q18 순환 불일치")
check(not any(r["source_id"] == r["target_id"] for r in seed), "자기 참조 관계가 저장됨")
path = gq["Q17"]["expected"]["path"]
ids = {v: k for k, v in name.items()}
check(all(ids[b] in calls[ids[a]] for a, b in zip(path, path[1:])), "Q17 경로에 없는 관계 포함")

# ---------------------------------------------------------------- 5. 파일 형식 검사
try:
    import yaml  # noqa: PLC0415
    y = yaml.safe_load((RAW / "02_개발_IT운영.yaml").read_text(encoding="utf-8"))
    check(len(y["agents"]) == 10, "YAML Agent 수")
except ImportError:
    pass
try:
    import openpyxl  # noqa: PLC0415
    ws = openpyxl.load_workbook(RAW / "04_인사_재무.xlsx")["Agent 현황"]
    check(ws.max_row == 11 and ws["B2"].value == "채용 서류 1차 검토 Agent", "XLSX 내용")
except ImportError:
    pass
if shutil.which("textutil"):
    out = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(RAW / "05_법무_경영기획.docx")],
                         capture_output=True, text=True)
    check(out.returncode == 0 and "계약서 검토 Agent" in out.stdout, "DOCX를 textutil로 열 수 없음")
ttl = (ROOT / "vocabulary/axiom_vocabulary.ttl").read_text(encoding="utf-8")
blocks = [b for b in re.split(r"\n\s*\n", ttl) if b.strip()]
for b in blocks:
    stmts = [s for s in b.split(" .\n") if s.strip() and not s.strip().startswith("@prefix") and not s.strip().startswith("#")]
    check(b.rstrip().endswith(".") or b.strip().startswith("#"), f"TTL 문장 끝 '.' 누락: {b[:40]!r}")

print(f"근거 위치 {n_ev}건, 관계 {len(rels)}건, 질의 {len(gold['queries'])}건 검사")
if problems:
    print("문제", len(problems), "건")
    for p in problems:
        print(" -", p)
    sys.exit(1)
print("OK — 원본 파일·어휘·정답 세트·질의 정답이 서로 맞습니다.")
